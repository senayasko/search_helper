"""ÖSYM'nin 2026 YKS yerleştirme tablolarından uygulama CSV'sini üretir.

Gereksinim: pip install openpyxl
"""
import csv
from pathlib import Path
from urllib.request import urlretrieve

from openpyxl import load_workbook


SOURCES = {
    "Lisans": "https://cdn.osym.gov.tr/en-kucuk-ve-en-buyuk-puanlar-tablo-4-rps0lq-18092428.xlsx",
    "Ön Lisans": "https://cdn.osym.gov.tr/en-kucuk-ve-en-buyuk-puanlar-tablo-3-0wptq7-18092428.xlsx",
}
HEADERS = ["Program Kodu", "Universite", "Bolum", "Puan Turu", "Taban Puan",
           "Tavan Puan", "Kontenjan", "Yerlesen", "Ogrenim Duzeyi", "Universite Turu"]


def number(value):
    return "" if value in (None, "--") else str(value)


def main():
    rows = []
    for level, url in SOURCES.items():
        path = Path(__file__).with_name(f"osym_2026_{'lisans' if level == 'Lisans' else 'onlisans'}.xlsx")
        if not path.exists():
            urlretrieve(url, path)
        sheet = load_workbook(path, read_only=True, data_only=True).active
        for row in sheet.iter_rows(min_row=4, values_only=True):
            if not row[0] or not str(row[0]).isdigit() or not row[4]:
                continue
            rows.append([str(row[0]), str(row[2]).strip(), str(row[4]).strip(),
                         str(row[5]).strip(), number(row[8]), number(row[9]),
                         number(row[6]), number(row[7]), level, str(row[1]).strip()])
    codes = [row[0] for row in rows]
    if len(codes) != len(set(codes)):
        raise ValueError("Tekrarlanan program kodu var")
    if len(rows) < 15000:
        raise ValueError("Beklenenden az program: kaynak dosyalarını kontrol edin")
    output = Path(__file__).with_name("universite_verileri.csv")
    temporary = output.with_suffix(".tmp")
    with temporary.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.writer(file)
        writer.writerow(HEADERS)
        writer.writerows(rows)
    temporary.replace(output)
    print(f"{len(rows)} program yazıldı: {output}")


if __name__ == "__main__":
    main()
